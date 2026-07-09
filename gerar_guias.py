#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera dois guias de apoio à candidatura de Jeferson dos Santos Cardoso:
1. Roteiro_de_Entrevista.docx  — pitch pessoal, pitch dos projetos (método STAR),
   respostas prontas para perguntas difíceis e mini-glossário técnico.
2. Plano_de_Estudos.docx — plano de 8 semanas até o início da Facens (ago/2026).

Mesma identidade visual do currículo: Arial, #333333, títulos #1B365D.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEXTO = RGBColor(0x33, 0x33, 0x33)
TITULO = RGBColor(0x1B, 0x36, 0x5D)
FONTE = "Arial"


def formatar_run(run, size, bold=False, italic=False, color=TEXTO):
    run.font.name = FONTE
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONTE)
    rFonts.set(qn("w:hAnsi"), FONTE)
    rFonts.set(qn("w:cs"), FONTE)


def paragrafo(doc, before=0, after=4, indent=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    return p


def texto(doc, partes, before=0, after=4, indent=None, size=11):
    p = paragrafo(doc, before=before, after=after, indent=indent)
    if isinstance(partes, str):
        partes = [(partes, False, False)]
    for t, b, i in partes:
        formatar_run(p.add_run(t), size, bold=b, italic=i)
    return p


def bullet(doc, conteudo, after=2, indent=0.5):
    p = paragrafo(doc, before=0, after=after)
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(-0.5)
    if isinstance(conteudo, str):
        conteudo = [(conteudo, False, False)]
    formatar_run(p.add_run("•\t"), 11)
    for t, b, i in conteudo:
        formatar_run(p.add_run(t), 11, bold=b, italic=i)
    return p


def fala(doc, texto_fala, after=6):
    """Bloco de fala pronta para decorar (itálico, recuado)."""
    p = paragrafo(doc, before=2, after=after, indent=0.75)
    formatar_run(p.add_run("“" + texto_fala + "”"), 11, italic=True)
    return p


def titulo_secao(doc, nome, before=12):
    p = paragrafo(doc, before=before, after=5)
    formatar_run(p.add_run(nome.upper()), 14, bold=True, color=TITULO)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1B365D")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def subtitulo(doc, nome, before=6):
    return texto(doc, [(nome, True, False)], before=before, after=3)


def novo_doc(titulo_pagina, subtitulo_pagina):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
    p = paragrafo(doc, before=0, after=1)
    formatar_run(p.add_run(titulo_pagina), 20, bold=True, color=TITULO)
    texto(doc, [(subtitulo_pagina, False, True)], after=4)
    return doc


# ═════════════════════════════════════════════════════════════════════════════
# DOCUMENTO 1 — ROTEIRO DE ENTREVISTA
# ═════════════════════════════════════════════════════════════════════════════
doc = novo_doc("ROTEIRO DE ENTREVISTA",
               "Jeferson dos Santos Cardoso — pitch pessoal, pitch dos projetos e "
               "respostas prontas. Treine em voz alta, gravando no celular.")

titulo_secao(doc, "1. Pitch pessoal (30–45 segundos)", before=6)
texto(doc, "Use quando pedirem “fale sobre você”. Decore a estrutura, não as palavras "
           "exatas — tem que soar natural.")
fala(doc, "Eu venho do chão de fábrica: comecei na logística, operando empilhadeira em "
          "câmara fria a -25 graus, e fui promovido de Júnior a Pleno em menos de um ano. "
          "Depois entrei no programa de aprendizagem da CBA, trabalhando de manhã no "
          "Suporte de Operações da Sala de Fornos 7 e estudando Técnico em Administração à "
          "tarde. Lá eu descobri minha vocação: percebi que os controles manuais do setor "
          "podiam virar sistemas, e criei por conta própria uma aplicação web de gestão de "
          "ocorrências que chegou a rodar em piloto. Isso me convenceu a seguir carreira em "
          "tecnologia — passei no Tecnólogo em Análise e Desenvolvimento de Sistemas da "
          "Facens e hoje busco uma oportunidade onde eu possa unir minha visão de operação "
          "com desenvolvimento de soluções.")

titulo_secao(doc, "2. Pitch do projeto principal (método STAR)")
texto(doc, "Sistema de Gestão de Ocorrências de Fornos — conte sempre nesta ordem: "
           "Situação, Tarefa, Ação, Resultado. Duração ideal: 2 a 3 minutos.")

subtitulo(doc, "SITUAÇÃO — a dor real")
fala(doc, "Na Sala de Fornos 7, as ocorrências de exaustão dos fornos eram anotadas em "
          "lousa. No fim do mês a lousa era apagada e todo o histórico se perdia. Ninguém "
          "conseguia responder perguntas simples como: qual forno dá mais problema? Qual "
          "tipo de falha é mais comum? A gestão não tinha dados para agir.")

subtitulo(doc, "TAREFA — o que você se propôs a fazer")
fala(doc, "Ninguém me pediu — eu decidi resolver. Meu objetivo era transformar um controle "
          "manual e efêmero em algo digital, auditável e visual, que qualquer operador "
          "conseguisse usar no celular.")

subtitulo(doc, "AÇÃO — como você fez (seja específico)")
bullet(doc, "Levantei os requisitos observando a rotina real: 4 categorias de ocorrência, "
            "cerca de 200 fornos, 4 turnos, operadores com pouco tempo para registrar.")
bullet(doc, "Usei IA (Claude, Copilot) como aceleradora: eu descrevia o que precisava, "
            "estudava o código gerado, testava, quebrava, corrigia e iterava até funcionar.")
bullet(doc, "Construí uma aplicação web (PWA) com dashboards: gráfico de Pareto por "
            "categoria, ranking dos fornos mais críticos, alertas automáticos e um modo TV "
            "para monitoramento na sala.")
bullet(doc, "Integrei com Power Automate: cada registro disparava uma requisição HTTP que "
            "gravava a linha numa planilha do Excel Online no OneDrive — espelho dos dados "
            "para a gestão.")
bullet(doc, "Publiquei via GitHub e Cloudflare e testei em piloto na rotina do setor.",
       after=4)

subtitulo(doc, "RESULTADO — e a lição de maturidade")
fala(doc, "O piloto funcionou: os registros passaram a gerar histórico e dava para ver na "
          "hora qual forno estava crítico. Depois, a TI da empresa descontinuou o sistema "
          "por política interna — aplicação não homologada e uso de marca. Foi uma aula de "
          "governança corporativa: hoje eu entendo por que empresas controlam shadow IT, e "
          "levaria o projeto primeiro à TI para homologação. Mas a experiência de construir, "
          "integrar e ver funcionando ninguém me tira.")

titulo_secao(doc, "3. Pitch do segundo projeto (1 minuto)")
texto(doc, "Sistema de Gestão de Estoque LIS — projeto acadêmico do SENAI.")
fala(doc, "Na disciplina de Logística Integrada, criei um site em HTML e JavaScript, "
          "hospedado no GitHub Pages, onde qualquer pessoa registrava entradas e saídas de "
          "estoque. Cada lançamento era enviado via Power Automate para uma planilha do "
          "Excel Online, que calculava saldo, estoque de segurança e ponto de pedido, e "
          "atualizava gráficos automaticamente — incluindo curva ABC. Foi onde aprendi na "
          "prática como front-end conversa com back-end através de uma API.")

titulo_secao(doc, "4. Perguntas difíceis — respostas prontas")

subtitulo(doc, "“Você fez esses projetos sozinho? Usou IA?”")
fala(doc, "Usei IA como ferramenta, sim — e considero isso um ponto forte. Eu descrevia o "
          "requisito, analisava o código gerado, testava e ajustava até funcionar. A IA "
          "acelera, mas não substitui: quem entendeu o problema, definiu os requisitos, "
          "testou com usuários reais e fez a integração funcionar fui eu. É assim que o "
          "mercado trabalha hoje, e eu aprendi a fazer isso cedo.")

subtitulo(doc, "“Por que o sistema saiu do ar?”")
fala(doc, "Decisão correta da TI da empresa: era uma aplicação não homologada usando o nome "
          "da companhia. Eu era aprendiz e não conhecia o processo de homologação. A lição "
          "que levo é envolver a TI e a governança desde o início — o que, aliás, me tornou "
          "mais preparado para trabalhar em ambiente corporativo.")

subtitulo(doc, "“O que é um PWA?”")
fala(doc, "Progressive Web App: um site que se comporta como aplicativo — dá para instalar "
          "no celular, tem ícone próprio e pode funcionar offline. Escolhi PWA porque os "
          "operadores usariam no celular sem precisar de loja de aplicativos.")

subtitulo(doc, "“Como funciona a integração com Power Automate?”")
fala(doc, "O Power Automate cria um fluxo com gatilho ‘quando uma solicitação HTTP é "
          "recebida’, que gera uma URL — um endpoint. Meu site enviava os dados do registro "
          "em JSON, com fetch do JavaScript, para essa URL. O fluxo recebia e adicionava uma "
          "linha na tabela do Excel Online. Na prática, é o conceito de API e webhook.")

subtitulo(doc, "“Qual sua experiência com dados?”")
fala(doc, "Excel Avançado e Power BI com uso diário real: na CBA eu monitorava dashboards "
          "de indicadores de segurança e tratava pendências. Nos meus projetos, criei "
          "dashboards de Pareto e ranking de criticidade. Sei que dado bom nasce de registro "
          "bem estruturado — uma linha por evento, sem célula mesclada — e levo isso para "
          "qualquer planilha ou sistema que eu construa.")

subtitulo(doc, "“Onde você se vê em 5 anos?”")
fala(doc, "Formado pela Facens, atuando como desenvolvedor ou analista, e sendo a ponte "
          "entre a operação e a tecnologia — porque eu falo as duas línguas: já vivi o chão "
          "de fábrica e sei programar a solução.")

titulo_secao(doc, "5. Mini-glossário (domine antes da entrevista)")
bullet(doc, [("API / Endpoint: ", True, False),
             ("porta de comunicação entre sistemas; o endereço (URL) que recebe ou entrega "
              "dados.", False, False)])
bullet(doc, [("Webhook: ", True, False),
             ("URL que fica “escutando”: quando um evento acontece, os dados são enviados "
              "para ela automaticamente.", False, False)])
bullet(doc, [("JSON: ", True, False),
             ("formato de texto para trocar dados entre sistemas (pares de campo e valor).",
              False, False)])
bullet(doc, [("Deploy: ", True, False),
             ("publicar a aplicação para o mundo usar (no seu caso: GitHub Pages e "
              "Cloudflare).", False, False)])
bullet(doc, [("Front-end / Back-end: ", True, False),
             ("o que o usuário vê (HTML/JS) versus o que processa por trás (fluxos, banco, "
              "planilha).", False, False)])
bullet(doc, [("Pareto (80/20): ", True, False),
             ("poucos itens causam a maioria dos problemas; o gráfico ordena as causas para "
              "priorizar ação.", False, False)])
bullet(doc, [("Curva ABC: ", True, False),
             ("classificação de itens de estoque por valor: A = poucos itens, muito valor; "
              "C = muitos itens, pouco valor.", False, False)])
bullet(doc, [("MTBF / MTTR: ", True, False),
             ("tempo médio entre falhas / tempo médio de reparo — indicadores de "
              "manutenção.", False, False)])
bullet(doc, [("Shadow IT: ", True, False),
             ("sistemas criados fora do controle da TI oficial — por isso seu piloto foi "
              "descontinuado; saber o termo impressiona.", False, False)], after=4)

titulo_secao(doc, "6. Regras de ouro no dia")
bullet(doc, "Fale do PROBLEMA e do RESULTADO antes da tecnologia — recrutador entende dor "
            "e valor, nem sempre entende código.")
bullet(doc, "Nunca minta nem infle: tudo neste roteiro aconteceu de verdade. Confiança vem "
            "da verdade bem contada.")
bullet(doc, "Treine cada pitch em voz alta 5 vezes, gravando no celular. Ouça e repita até "
            "soar natural.")
bullet(doc, "Leve o portfólio: currículo impresso + GitHub público com README dos projetos.")
bullet(doc, "Ao final, pergunte algo: “Como é o dia a dia do time?” ou “O que vocês esperam "
            "de alguém nos primeiros 90 dias?” — demonstra interesse genuíno.", after=0)

doc.save("Roteiro_de_Entrevista.docx")
print("OK: Roteiro_de_Entrevista.docx")

# ═════════════════════════════════════════════════════════════════════════════
# DOCUMENTO 2 — PLANO DE ESTUDOS
# ═════════════════════════════════════════════════════════════════════════════
doc = novo_doc("PLANO DE ESTUDOS — 8 SEMANAS",
               "Jeferson dos Santos Cardoso — preparação até o início da Facens "
               "(agosto/2026). Meta: 1h30 por dia, 6 dias por semana.")

titulo_secao(doc, "Como usar este plano", before=6)
bullet(doc, "Cada semana tem UM foco e UMA entrega. A entrega é inegociável — é ela que "
            "vira portfólio.")
bullet(doc, "Use IA para aprender, não para copiar: peça explicação linha a linha do código "
            "que ela gerar e só avance quando entender.")
bullet(doc, "Tudo que você produzir vai para o seu GitHub pessoal (github.com/j3ff-11) com "
            "README explicando o projeto.", after=4)

titulo_secao(doc, "Semanas 1–2 — Lógica e Python de verdade")
texto(doc, [("Objetivo: ", True, False),
            ("sair do “fundamentos” para escrever programas úteis sozinho.", False, False)])
bullet(doc, "Curso gratuito: “Curso em Vídeo — Python (Gustavo Guanabara)”, mundos 1 e 2, "
            "no YouTube.")
bullet(doc, "Prática diária: 3 exercícios no site “Beecrowd” (categoria iniciante).")
bullet(doc, [("Entrega: ", True, False),
             ("script Python que lê uma planilha de ocorrências (CSV) e imprime o ranking "
              "dos fornos mais críticos — a versão em código do que você fazia na lousa.",
              False, False)], after=4)

titulo_secao(doc, "Semanas 3–4 — HTML, CSS e JavaScript")
texto(doc, [("Objetivo: ", True, False),
            ("entender de verdade o que a IA gerava nos seus sites.", False, False)])
bullet(doc, "Curso gratuito: freeCodeCamp “Responsive Web Design” (em inglês, com legenda) "
            "ou Curso em Vídeo — HTML5/CSS3.")
bullet(doc, "JavaScript: foque em manipular formulários, eventos e fetch (foi exatamente o "
            "que seu sistema LIS usava).")
bullet(doc, [("Entrega: ", True, False),
             ("reconstruir o Sistema LIS do zero, com nome neutro e dados fictícios, "
              "publicado no SEU GitHub Pages — sem marca de empresa. Vira o link vivo do "
              "seu currículo.", False, False)], after=4)

titulo_secao(doc, "Semanas 5–6 — Git, GitHub e portfólio profissional")
texto(doc, [("Objetivo: ", True, False),
            ("dominar a ferramenta que todo dev usa todo dia.", False, False)])
bullet(doc, "Aprenda na prática: git init, add, commit, push, branch, merge — usando os "
            "próprios projetos das semanas anteriores.")
bullet(doc, "Escreva um README caprichado para cada repositório: problema, solução, "
            "tecnologias, prints das telas.")
bullet(doc, [("Entrega: ", True, False),
             ("perfil GitHub arrumado com 2–3 repositórios documentados e um README de "
              "perfil se apresentando.", False, False)], after=4)

titulo_secao(doc, "Semanas 7–8 — SQL e reconstrução do projeto principal")
texto(doc, [("Objetivo: ", True, False),
            ("aprender a linguagem de dados que a Facens vai cobrar e fechar o portfólio "
             "com chave de ouro.", False, False)])
bullet(doc, "Curso gratuito: “SQL para Análise de Dados” (YouTube — Curso em Vídeo MySQL "
            "ou freeCodeCamp SQL).")
bullet(doc, "Pratique com os dados fictícios de ocorrências: SELECT, WHERE, GROUP BY, "
            "ORDER BY — as mesmas perguntas do Pareto, agora em SQL.")
bullet(doc, [("Entrega: ", True, False),
             ("versão portfólio do sistema de ocorrências (nome neutro, dados fictícios), "
              "publicada no seu GitHub Pages com dashboards funcionando.", False, False)],
       after=4)

titulo_secao(doc, "Contínuo (todas as semanas)")
bullet(doc, "Inglês técnico: 15 min/dia lendo documentação ou vendo vídeo com legenda — "
            "80% do material bom de programação está em inglês.")
bullet(doc, "1 pitch de entrevista treinado em voz alta por semana (use o Roteiro de "
            "Entrevista).")
bullet(doc, "LinkedIn ativo: 1 post curto por semana contando o que aprendeu/construiu — "
            "recrutador júnior olha consistência, não perfeição.", after=4)

titulo_secao(doc, "Depois das 8 semanas (Facens começa)")
bullet(doc, "Priorize as disciplinas da faculdade — o plano acima te coloca À FRENTE da "
            "turma, não substitui o curso.")
bullet(doc, "Candidate-se a estágios de TI desde o 1º semestre: com portfólio no GitHub e "
            "sua história de operação, você compete de igual para igual.")
bullet(doc, "Meta de 12 meses: primeiro estágio ou júnior em desenvolvimento, dados ou "
            "automação.", after=0)

doc.save("Plano_de_Estudos.docx")
print("OK: Plano_de_Estudos.docx")
