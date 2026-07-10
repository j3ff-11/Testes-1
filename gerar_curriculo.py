#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador do currículo de Jeferson dos Santos Cardoso — versão única 10/10.

Especificação visual:
- Fonte: Arial | Nome: 24 pt negrito | Títulos: 14 pt negrito MAIÚSCULAS | Corpo: 11 pt
- Espaçamento entre linhas: 1,15 | Margens: 2,0 cm | Alinhado à esquerda
- Cores: fundo #FFFFFF, texto #333333, títulos #1B365D
- Linha horizontal fina (#1B365D) abaixo de cada título de seção
- Sem ícones, gráficos, colunas, tabelas ou caixas de texto | 2 páginas exatas
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEXTO = RGBColor(0x33, 0x33, 0x33)
TITULO = RGBColor(0x1B, 0x36, 0x5D)
FONTE = "Arial"


def formatar_run(run, size, bold=False, italic=False, color=TEXTO):
    run.font.name = FONTE
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONTE)
    rFonts.set(qn("w:hAnsi"), FONTE)
    rFonts.set(qn("w:cs"), FONTE)


def paragrafo(doc, before=0, after=4, indent=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    return p


def texto(doc, partes, before=0, after=3, indent=None, size=11):
    p = paragrafo(doc, before=before, after=after, indent=indent)
    if isinstance(partes, str):
        partes = [(partes, False, False)]
    for t, b, i in partes:
        formatar_run(p.add_run(t), size, bold=b, italic=i)
    return p


def bullet(doc, conteudo, after=1):
    p = paragrafo(doc, before=0, after=after)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.first_line_indent = Cm(-0.5)
    if isinstance(conteudo, str):
        conteudo = [(conteudo, False, False)]
    formatar_run(p.add_run("•\t"), 11)
    for t, b, i in conteudo:
        formatar_run(p.add_run(t), 11, bold=b, italic=i)
    return p


def titulo_secao(doc, nome, before=7):
    p = paragrafo(doc, before=before, after=4)
    p.paragraph_format.keep_with_next = True
    formatar_run(p.add_run(nome.upper()), 14, bold=True, color=TITULO)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1B365D")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# ── Cabeçalho ────────────────────────────────────────────────────────────────
p = paragrafo(doc, before=0, after=2)
formatar_run(p.add_run("JEFERSON DOS SANTOS CARDOSO"), 24, bold=True, color=TITULO)

texto(doc, "Mairinque – SP  |  (11) 95640-3941  |  jeferson20sc@gmail.com", after=1)
texto(doc, "linkedin.com/in/jeferson-santos35  |  github.com/j3ff-11  |  CNH: A/B (EAR)",
      after=1)
texto(doc, [("Suporte de Operações | Administração | BI | Tecnologia | Logística",
             True, False)], after=2)

# ── Resumo profissional ──────────────────────────────────────────────────────
titulo_secao(doc, "Resumo Profissional", before=6)
texto(doc,
      "Profissional com formação técnica em Eletromecânica e Administração pelo SENAI, "
      "aprovado no Tecnólogo em Análise e Desenvolvimento de Sistemas (Facens), com "
      "experiência em ambiente industrial, logística, operação, organização de processos, "
      "registros, indicadores e melhoria contínua. Atuação na CBA com suporte às rotinas "
      "operacionais da Sala de Fornos 7, governança documental, acompanhamento de "
      "indicadores, controle de almoxarifado, SAP, Power BI, Power Apps e 5S — além do "
      "desenvolvimento autoral de um sistema web de gestão de ocorrências, testado em "
      "piloto no setor. Experiência anterior como Operador de Armazém Júnior e "
      "Pleno, com operação de empilhadeira, WMS, expedição e carregamento em ambiente de "
      "alto giro. Perfil disciplinado, responsável, hands-on e com facilidade para aprender "
      "procedimentos e cumprir normas de segurança, qualidade e produtividade.",
      after=3)

# ── Experiência profissional ─────────────────────────────────────────────────
titulo_secao(doc, "Experiência Profissional")

texto(doc, [("COMPANHIA BRASILEIRA DE ALUMÍNIO — CBA", True, False),
            ("  |  Alumínio/SP", False, False)], before=1, after=1)
texto(doc, [("Aprendiz Técnico em Administração – Suporte de Operações Industriais — "
             "Sala de Fornos 7", True, False),
            ("  |  ", False, False), ("jan/2025 – jun/2026", False, True)], after=2)
bullet(doc, "Gestão e governança documental no Docnix: pesquisa avançada, verificação de "
            "vigência e distribuição de Procedimentos Operacionais, Padrões de Trabalho e "
            "FMEAs.")
bullet(doc, "Apoio local de TI/Key User: chamados no ServiceNow, suporte a acessos, "
            "impressoras e equipamentos bloqueados por BitLocker.")
bullet(doc, "Monitoramento diário de indicadores de segurança e operação (DDS, "
            "abrangências e pendências) em Power BI, Power Apps e ObraSoft.")
bullet(doc, "Planilhas de controle em Excel para inventários, entrada e saída de materiais "
            "e distribuição de EPIs, com consultas de saldo no SAP e organização do "
            "almoxarifado com 5S.")
bullet(doc, "Identificação de falhas em controles manuais e criação de checklists e "
            "controles digitais que evoluíram para um sistema web de gestão de "
            "ocorrências, testado em piloto na área.", after=4)

texto(doc, [("CEFRI / SUPERFRIO — Logística e Armazenagem Frigorificada", True, False),
            ("  |  Mairinque/SP", False, False)], after=1)
texto(doc, [("Operador de Armazém Júnior → Pleno", True, False),
            ("  |  ", False, False), ("abr/2023 – jan/2025", False, True)], after=2)
bullet(doc, "Promovido de Operador Júnior a Pleno em menos de 12 meses, em reconhecimento "
            "ao desempenho, confiabilidade e produtividade.")
bullet(doc, "Operação de empilhadeiras (NR-11) em câmaras frias de até -25 °C, com "
            "picking, conferência, armazenagem e carregamento em turno noturno de alta "
            "demanda.")
bullet(doc, "Utilização de coletor de dados e WMS Blue Yonder para controle de "
            "movimentações, separação de pedidos e apoio ao fluxo de estoque.")
bullet(doc, "Capacitação contínua em paralelo ao trabalho: Excel Avançado, Power BI, "
            "Fundamentos de Python, PCP e Metrologia.", after=4)

texto(doc, [("SUPERMERCADO SÃO ROQUE — Centro de Distribuição", True, False),
            ("  |  São Roque/SP", False, False)], after=1)
texto(doc, [("Ajudante Operacional", True, False),
            ("  |  ", False, False), ("ago/2021 – jul/2022", False, True)], after=2)
bullet(doc, "Carregamento de caminhões, paletização, organização de estoque, separação de "
            "cargas e devolução de ativos retornáveis.", after=3)

# ── Projetos de tecnologia ───────────────────────────────────────────────────
titulo_secao(doc, "Projetos de Tecnologia")

texto(doc, [("Sistema de Gestão de Ocorrências de Fornos Industriais — Projeto Autoral",
             True, False)], before=1, after=2)
bullet(doc, "Dor real identificada na operação: controles manuais em lousa, baixa "
            "rastreabilidade e ausência de histórico estruturado para análise de "
            "recorrências.")
bullet(doc, "Desenvolvimento de aplicação web (PWA) para registrar ocorrências, "
            "acompanhar histórico, visualizar fornos críticos, gerar rankings e aplicar "
            "análise de Pareto.")
bullet(doc, "Integração com Power Automate e Excel Online, deploy via GitHub/Cloudflare e "
            "desenvolvimento assistido por IA (Claude, Copilot e ChatGPT).")
bullet(doc, [("Resultado: ", True, False),
             ("substituição de controles soltos por um modelo digital rastreável e "
              "visual, testado em piloto, facilitando a priorização de problemas.",
              False, False)], after=4)

texto(doc, [("Sistema de Gestão de Estoque LIS — Projeto Acadêmico SENAI", True, False)],
      after=2)
bullet(doc, "Site em HTML/JavaScript (GitHub Pages) integrado ao Excel Online via Power "
            "Automate: curva ABC, estoque de segurança, ponto de pedido e dashboards em "
            "tempo real.", after=3)

# ── Formação acadêmica ───────────────────────────────────────────────────────
titulo_secao(doc, "Formação Acadêmica")

texto(doc, [("Tecnólogo em Análise e Desenvolvimento de Sistemas — Facens", True, False),
            ("  |  início em agosto/2026", False, False)], before=1, after=2)
texto(doc, [("Técnico em Administração — SENAI", True, False),
            ("  |  1.200h — conclusão: junho/2026", False, False)], after=2)
texto(doc, [("Técnico em Eletromecânica — SENAI Mairinque", True, False),
            ("  |  1.500h — conclusão: junho/2025", False, False)], after=3)

# ── Cursos e certificações ───────────────────────────────────────────────────
titulo_secao(doc, "Cursos e Certificações")

bullet(doc, "NR-11 — Operação de Empilhadeira — SENAI (32h, 2023); reciclagens — "
            "SuperFrio (2024)")
bullet(doc, "Metrologia Aplicada à Mecânica (60h) e PCP (40h) — SENAI (2024)")
bullet(doc, "Caldeiraria Prática — AHCX Treinamentos / MJS Brasil (jan–jul/2026, em "
            "conclusão)")
bullet(doc, "Excel Avançado (40h), Power BI (32h) e Fundamentos do Python (30h) — SENAI "
            "(2024)")
bullet(doc, "Lógica de Programação (14h), Segurança Cibernética (4h), Indústria 4.0 "
            "(20h), Fundamentos da IA (8h) e Ética na IA (4h) — SENAI (2025)", after=3)

# ── Competências técnicas ────────────────────────────────────────────────────
titulo_secao(doc, "Competências Técnicas")

bullet(doc, [("Tecnologia, Dados e Automação: ", True, False),
             ("Excel Avançado, Power BI, Power Automate, Power Apps, HTML, JavaScript, "
              "GitHub, fundamentos de Python e desenvolvimento assistido por IA.",
              False, False)])
bullet(doc, [("Sistemas Corporativos: ", True, False),
             ("ServiceNow, Docnix, SAP, WMS Blue Yonder e ObraSoft.", False, False)])
bullet(doc, [("Administração e Operações: ", True, False),
             ("governança documental, relatórios, indicadores, inventários, controle de "
              "materiais, gestão de EPIs e organização de almoxarifado.", False, False)])
bullet(doc, [("Logística e Indústria: ", True, False),
             ("operação de empilhadeira (NR-11), expedição, armazenagem, picking, câmaras "
              "frias, PCP, metrologia, 5S, PDCA, Matriz GUT e gestão visual.",
              False, False)])
bullet(doc, [("Comportamentais: ", True, False),
             ("disciplina, responsabilidade, proatividade, aprendizado rápido, resiliência "
              "e mentalidade de melhoria contínua.", False, False)], after=3)

# ── Informações adicionais ───────────────────────────────────────────────────
titulo_secao(doc, "Informações Adicionais")

bullet(doc, "CNH categorias A e B (EAR); disponibilidade para início imediato, em turnos "
            "ou horário comercial; interesse em suporte de TI, administração, BI, "
            "logística, operações e melhoria de processos.", after=0)

doc.save("Curriculo_Jeferson_dos_Santos_Cardoso.docx")
print("OK: Curriculo_Jeferson_dos_Santos_Cardoso.docx gerado.")
